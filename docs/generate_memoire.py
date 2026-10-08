# -*- coding: utf-8 -*-
"""Génère le mémoire Word ARSAVE (chapitres 1 à 4) et les figures UML."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent / "memoire"
OUT.mkdir(exist_ok=True)

NAVY = (15, 42, 92)
BLUE = (29, 78, 216)
LIGHT = (232, 241, 255)
LINE = (30, 41, 59)
WHITE = (255, 255, 255)
GREY = (71, 85, 105)
PALE = (248, 250, 252)


def font(size, bold=False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(rf"C:\Windows\Fonts\{name}", size)


def wrap(draw, text, fnt, max_w):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = w if not cur else cur + " " + w
        if draw.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def text_center(draw, cx, cy, text, fnt, fill=LINE, max_w=160):
    lines = wrap(draw, text, fnt, max_w)
    heights = []
    for line in lines:
        box = draw.textbbox((0, 0), line, font=fnt)
        heights.append(box[3] - box[1])
    total = sum(heights) + 2 * (len(lines) - 1)
    y = cy - total / 2
    for line, h in zip(lines, heights):
        w = draw.textlength(line, font=fnt)
        draw.text((cx - w / 2, y), line, font=fnt, fill=fill)
        y += h + 2


def arrow(draw, x1, y1, x2, y2, color=LINE, width=2, dashed=False):
    if dashed:
        import math
        length = math.hypot(x2 - x1, y2 - y1) or 1
        ux, uy = (x2 - x1) / length, (y2 - y1) / length
        pos = 0
        on = True
        while pos < length - 10:
            n = min(8, length - 10 - pos)
            if on:
                draw.line(
                    [(x1 + ux * pos, y1 + uy * pos), (x1 + ux * (pos + n), y1 + uy * (pos + n))],
                    fill=color,
                    width=width,
                )
            pos += n
            on = not on
    else:
        draw.line([(x1, y1), (x2, y2)], fill=color, width=width)
    import math
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 10
    p1 = (x2 - size * math.cos(ang - 0.4), y2 - size * math.sin(ang - 0.4))
    p2 = (x2 - size * math.cos(ang + 0.4), y2 - size * math.sin(ang + 0.4))
    draw.polygon([(x2, y2), p1, p2], fill=color)


def actor(draw, x, y, name):
    draw.ellipse((x - 16, y - 58, x + 16, y - 26), outline=NAVY, width=2)
    draw.line((x, y - 26, x, y + 18), fill=NAVY, width=2)
    draw.line((x - 22, y - 8, x + 22, y - 8), fill=NAVY, width=2)
    draw.line((x, y + 18, x - 18, y + 48), fill=NAVY, width=2)
    draw.line((x, y + 18, x + 18, y + 48), fill=NAVY, width=2)
    text_center(draw, x, y + 68, name, font(16, True), NAVY, 120)


def ellipse(draw, cx, cy, w, h, text):
    draw.ellipse((cx - w / 2, cy - h / 2, cx + w / 2, cy + h / 2), fill=LIGHT, outline=BLUE, width=2)
    text_center(draw, cx, cy, text, font(15), LINE, w - 28)


def box(draw, x, y, w, h, title, lines, fill=WHITE):
    draw.rounded_rectangle((x, y, x + w, y + h), radius=6, fill=fill, outline=NAVY, width=2)
    draw.rectangle((x, y, x + w, y + 28), fill=NAVY)
    text_center(draw, x + w / 2, y + 14, title, font(14, True), WHITE, w - 12)
    yy = y + 40
    for line in lines:
        draw.text((x + 10, yy), line, font=font(13), fill=LINE)
        yy += 18


def save(img, name):
    path = OUT / name
    img.save(path, "PNG")
    return path


def fig_usecase():
    img = Image.new("RGB", (1500, 980), WHITE)
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((250, 30, 1180, 940), radius=12, outline=NAVY, width=3)
    d.text((270, 42), "Système ARSAVE", font=font(18, True), fill=NAVY)
    actor(d, 120, 430, "Utilisateur")
    actor(d, 1360, 620, "Administrateur")
    cases = [
        (470, 150, "S'inscrire"),
        (760, 150, "S'authentifier"),
        (1050, 150, "Consulter le tableau de bord"),
        (470, 340, "Sauvegarder un fichier"),
        (760, 340, "Restaurer un fichier"),
        (1050, 340, "Consulter l'historique"),
        (470, 540, "Mettre en corbeille"),
        (760, 540, "Gérer le profil"),
        (1050, 700, "Gérer les comptes"),
        (760, 760, "Chiffrer le fichier"),
    ]
    for x, y, t in cases:
        ellipse(d, x, y, 250, 90, t)
    targets = [
        (470, 150), (760, 150), (1050, 150),
        (470, 340), (760, 340), (1050, 340),
        (470, 540), (760, 540),
    ]
    for x, y in targets:
        d.line((175, 400, x - 125, y), fill=GREY, width=2)
    d.line((1320, 620, 1175, 700), fill=GREY, width=2)
    d.line((595, 340, 660, 720), fill=BLUE, width=2)
    d.text((610, 500), "<<include>>", font=font(13, True), fill=BLUE)
    return save(img, "fig_cas_utilisation.png")


def lifelines(title, names, messages, height=980):
    img = Image.new("RGB", (1500, height), WHITE)
    d = ImageDraw.Draw(img)
    d.text((40, 16), title, font=font(18, True), fill=NAVY)
    n = len(names)
    xs = [180 + i * (1180 / max(n - 1, 1)) for i in range(n)]
    top = 90
    for x, name in zip(xs, names):
        tw = d.textlength(name, font=font(14, True))
        d.rounded_rectangle((x - tw / 2 - 12, top, x + tw / 2 + 12, top + 36), radius=6, fill=LIGHT, outline=NAVY, width=2)
        text_center(d, x, top + 18, name, font(14, True), NAVY, tw + 8)
        y = top + 36
        while y < height - 40:
            d.line((x, y, x, min(y + 10, height - 40)), fill=GREY, width=1)
            y += 16
    y = 180
    for src, dst, label, dashed in messages:
        x1, x2 = xs[src], xs[dst]
        color = GREY if dashed else BLUE
        if src == dst:
            d.line([(x1, y), (x1 + 90, y)], fill=color, width=2)
            d.line([(x1 + 90, y), (x1 + 90, y + 36)], fill=color, width=2)
            arrow(d, x1 + 90, y + 36, x1 + 2, y + 36, color, 2, False)
            d.text((x1 + 98, y + 6), label, font=font(13), fill=LINE)
            y += 78
        else:
            arrow(d, x1, y, x2, y, color, 2, dashed)
            text_center(d, (x1 + x2) / 2, y - 16, label, font(13), LINE, max(abs(x2 - x1) - 30, 80))
            y += 62
    return img


def fig_seq_auth():
    img = lifelines(
        "Scénario nominal : authentification",
        ["Utilisateur", "Interface", "API Auth", "MySQL"],
        [
            (0, 1, "1. Saisit email et mot de passe", False),
            (1, 2, "2. POST /auth/login", False),
            (2, 3, "3. SELECT utilisateur", False),
            (3, 2, "4. hash, sel KDF", True),
            (2, 2, "5. Vérifie le mot de passe", False),
            (2, 3, "6. INSERT session (jeton haché)", False),
            (2, 1, "7. jeton, sel, profil", True),
            (1, 1, "8. Dérive la clé maître (PBKDF2)", False),
            (1, 0, "9. Accès au tableau de bord", True),
        ],
        820,
    )
    return save(img, "fig_sequence_auth.png")


def fig_seq_backup():
    img = lifelines(
        "Scénario nominal : sauvegarde d'un fichier",
        ["Utilisateur", "Interface", "Crypto", "API Fichiers", "Stockage"],
        [
            (0, 1, "1. Choisit un fichier", False),
            (1, 2, "2. Demande le chiffrement", False),
            (2, 2, "3. AES-256-GCM", False),
            (2, 1, "4. blob, nonce, clé enveloppée", True),
            (1, 3, "5. POST /files (multipart)", False),
            (3, 4, "6. Écrit le blob chiffré", False),
            (3, 3, "7. INSERT fichier et version", False),
            (3, 1, "8. Métadonnées enregistrées", True),
            (1, 0, "9. Confirmation à l'écran", True),
        ],
        820,
    )
    return save(img, "fig_sequence_sauvegarde.png")


def fig_activity():
    img = Image.new("RGB", (980, 1280), WHITE)
    d = ImageDraw.Draw(img)
    d.text((40, 16), "Activité : sauvegarder un fichier", font=font(18, True), fill=NAVY)

    def node(cx, cy, text, kind="action"):
        if kind == "start":
            d.ellipse((cx - 16, cy - 16, cx + 16, cy + 16), fill=NAVY)
        elif kind == "end":
            d.ellipse((cx - 18, cy - 18, cx + 18, cy + 18), outline=NAVY, width=3)
            d.ellipse((cx - 12, cy - 12, cx + 12, cy + 12), fill=NAVY)
        elif kind == "decision":
            d.polygon([(cx, cy - 36), (cx + 90, cy), (cx, cy + 36), (cx - 90, cy)], fill=LIGHT, outline=BLUE)
            text_center(d, cx, cy, text, font(13), LINE, 150)
        else:
            d.rounded_rectangle((cx - 150, cy - 28, cx + 150, cy + 28), radius=10, fill=LIGHT, outline=NAVY, width=2)
            text_center(d, cx, cy, text, font(15), LINE, 270)

    steps = [
        (490, 70, "", "start"),
        (490, 140, "Choisir un fichier", "action"),
        (490, 230, "Fichier lisible ?", "decision"),
        (490, 330, "Générer une clé de fichier", "action"),
        (490, 410, "Chiffrer AES-256-GCM", "action"),
        (490, 490, "Envelopper la clé avec la clé maître", "action"),
        (490, 580, "Envoyer le blob à l'API", "action"),
        (490, 680, "Enregistrement accepté ?", "decision"),
        (490, 780, "Afficher la confirmation", "action"),
        (490, 870, "", "end"),
    ]
    for cx, cy, text, kind in steps:
        node(cx, cy, text, kind)
    node(760, 230, "Afficher l'erreur", "action")
    node(760, 680, "Afficher l'échec", "action")
    for y1, y2 in ((86, 112), (168, 194), (266, 302), (358, 382), (438, 462), (518, 552), (608, 644), (716, 752), (808, 852)):
        arrow(d, 490, y1, 490, y2)
    arrow(d, 580, 230, 610, 230)
    d.text((588, 206), "non", font=font(13, True), fill=BLUE)
    arrow(d, 580, 680, 610, 680)
    d.text((588, 656), "non", font=font(13, True), fill=BLUE)
    d.text((450, 248), "oui", font=font(13, True), fill=BLUE)
    d.text((450, 698), "oui", font=font(13, True), fill=BLUE)
    return save(img, "fig_activite.png")


def fig_classes():
    img = Image.new("RGB", (1600, 980), WHITE)
    d = ImageDraw.Draw(img)
    d.text((40, 16), "Diagramme de classes du domaine ARSAVE", font=font(18, True), fill=NAVY)
    specs = [
        (620, 50, "Utilisateur", ["id", "email", "nom", "role", "hashMotDePasse", "selKdf"]),
        (80, 360, "Session", ["id", "jetonHache", "expireLe"]),
        (430, 360, "Appareil", ["id", "uid", "nom"]),
        (780, 360, "Fichier", ["id", "nomLogique", "mime", "statut"]),
        (1140, 360, "Operation", ["id", "type", "statut", "message", "creeLe"]),
        (780, 680, "VersionFichier", ["id", "numero", "taille", "chemin", "nonce", "cleEnveloppee"]),
    ]
    for x, y, title, lines in specs:
        h = 36 + 18 * len(lines) + 16
        box(d, x, y, 280, h, title, lines, PALE)
    arrow(d, 700, 200, 220, 360)
    d.text((400, 250), "1", font=font(14, True), fill=BLUE)
    d.text((230, 330), "*", font=font(14, True), fill=BLUE)
    arrow(d, 760, 200, 560, 360)
    d.text((640, 250), "1", font=font(14, True), fill=BLUE)
    d.text((540, 330), "*", font=font(14, True), fill=BLUE)
    arrow(d, 820, 200, 900, 360)
    d.text((860, 250), "1", font=font(14, True), fill=BLUE)
    d.text((880, 330), "*", font=font(14, True), fill=BLUE)
    arrow(d, 900, 200, 1200, 360)
    d.text((1080, 240), "1", font=font(14, True), fill=BLUE)
    d.text((1160, 330), "*", font=font(14, True), fill=BLUE)
    arrow(d, 920, 530, 920, 680)
    d.text((930, 590), "1", font=font(14, True), fill=BLUE)
    d.text((930, 650), "*", font=font(14, True), fill=BLUE)
    return save(img, "fig_classes.png")


def fig_architecture():
    img = Image.new("RGB", (1400, 520), WHITE)
    d = ImageDraw.Draw(img)
    items = [
        (60, 160, "Client Flutter\n(Web / Android)", "Chiffrement local\nAES-256-GCM"),
        (520, 160, "API REST PHP", "Auth, fichiers,\nadministration"),
        (980, 80, "MySQL", "Comptes, sessions,\nmétadonnées"),
        (980, 280, "Stockage chiffré", "Blobs uniquement"),
    ]
    for x, y, title, sub in items:
        d.rounded_rectangle((x, y, x + 320, y + 150), radius=14, fill=LIGHT, outline=NAVY, width=2)
        text_center(d, x + 160, y + 48, title, font(18, True), NAVY, 290)
        text_center(d, x + 160, y + 105, sub, font(15), GREY, 290)
    arrow(d, 380, 235, 520, 235)
    d.text((400, 200), "HTTPS + jeton", font=font(13, True), fill=BLUE)
    arrow(d, 840, 210, 980, 150)
    arrow(d, 840, 260, 980, 330)
    return save(img, "fig_architecture.png")


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=12, bold=False, italic=False, center=False, space_after=8, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    if first_line and not center:
        p.paragraph_format.first_line_indent = Cm(1.25)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h(doc, text, level):
    if level == 1 and not text.startswith("Chapitre 1"):
        doc.add_page_break()
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True, color=(15, 42, 92))
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_fig(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.5))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(12)
    r = c.add_run(caption)
    set_run_font(r, size=11, italic=True)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=(255, 255, 255))
        shading = cell._tc.get_or_add_tcPr()
        sh = shading.makeelement(qn("w:shd"), {qn("w:fill"): "0F2A5C", qn("w:val"): "clear"})
        shading.append(sh)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=10)
    doc.add_paragraph()


def build():
    figs = {
        "uc": fig_usecase(),
        "s1": fig_seq_auth(),
        "s2": fig_seq_backup(),
        "act": fig_activity(),
        "cls": fig_classes(),
        "arch": fig_architecture(),
    }
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_p(doc, "MÉMOIRE", size=14, bold=True, center=True, first_line=False, space_after=6)
    add_p(doc, "ARSAVE", size=26, bold=True, center=True, first_line=False, space_after=4)
    add_p(
        doc,
        "Application web de sauvegarde, de restauration et d'administration sécurisées de fichiers",
        size=14,
        italic=True,
        center=True,
        first_line=False,
        space_after=12,
    )
    add_p(doc, "Chapitre 1 — Revue de littérature", center=True, first_line=False, space_after=2)
    add_p(doc, "Chapitre 2 — Analyse de systèmes existants", center=True, first_line=False, space_after=2)
    add_p(doc, "Chapitre 3 — Analyse et conception du système proposé", center=True, first_line=False, space_after=2)
    add_p(doc, "Chapitre 4 — Implémentation du système", center=True, first_line=False, space_after=10)
    add_p(
        doc,
        "Conception conduite exclusivement selon la méthode Unified Process (UP) et la notation UML.",
        italic=True,
        center=True,
        first_line=False,
    )
    doc.add_page_break()

    add_h(doc, "Chapitre 1. Revue de littérature", 1)
    add_h(doc, "Introduction", 2)
    add_p(
        doc,
        "Ce chapitre pose le cadre théorique du mémoire. Il ne décrit pas encore les écrans d'ARSAVE : il rassemble les notions qui justifient le système, afin que la conception du chapitre 3 et la construction du chapitre 4 s'appuient sur des sources identifiées. Le fil conducteur est le suivant : une sauvegarde distante n'est utile que si l'utilisateur peut récupérer ses fichiers, et elle n'est acceptable, dans le périmètre retenu, que si l'hébergeur ne peut pas les lire.",
    )
    add_p(
        doc,
        "La revue croise quatre domaines. Le premier est le stockage distant et l'architecture REST (Fielding, 2000 ; Fielding et Reschke, 2014). Le deuxième est le chiffrement authentifié et la dérivation de clé à partir d'un mot de passe (NIST, 2001 ; Dworkin, 2007 ; Moriarty, Kaliski et Rusch, 2017 ; Turan, Barker, Burr et Chen, 2010). Le troisième est la qualité et la protection des données (ISO/IEC, 2011 ; ISO/IEC, 2022 ; OWASP Foundation, s. d.-a). Le quatrième est la méthode de conduite du projet : le Unified Process et UML (Jacobson, Booch et Rumbaugh, 1999 ; Larman, 2004 ; OMG, 2017). Les sections qui suivent développent ces points, puis une conclusion en tire les exigences transmises au chapitre 2.",
    )
    add_h(doc, "1.1 Objet de la revue", 2)
    add_p(
        doc,
        "La sauvegarde de fichiers personnels ou professionnels ne se limite plus à une copie locale. Elle suppose un service distant, une reprise après perte de l'appareil et une garantie que le contenu ne soit pas lisible par l'hébergeur. Cette revue situe ARSAVE par rapport à trois ensembles de travaux et de normes : le stockage et la sauvegarde, le chiffrement côté client, puis le cadre de conception retenu, le Unified Process accompagné d'UML (Jacobson, Booch et Rumbaugh, 1999 ; OMG, 2017).",
    )
    add_h(doc, "1.2 Sauvegarde et accès distant", 2)
    add_p(
        doc,
        "Un service de sauvegarde sépare en général trois responsabilités : l'identité de l'utilisateur, les métadonnées du fichier (nom, type, versions, état) et le contenu binaire. L'architecture REST, décrite par Fielding (2000), permet d'exposer ces responsabilités par des ressources et des verbes HTTP, sans coupler le client au moteur de base de données. Le transport est protégé par TLS (Rescorla, 2018), ce qui authentifie le canal mais ne chiffre pas le fichier vis-à-vis du serveur une fois la requête déchiffrée au niveau applicatif.",
    )
    add_p(
        doc,
        "La littérature opérationnelle des grands services (Google, s. d. ; Dropbox, s. d.) décrit surtout un chiffrement au repos maîtrisé par l'opérateur : le fournisseur détient les clés de stockage. Ce modèle facilite la recherche, la prévisualisation et le partage, mais il place le contenu en clair dans le périmètre de confiance du serveur. ARSAVE part du constat inverse : le serveur ne doit conserver que des métadonnées et un blob déjà chiffré.",
    )
    add_h(doc, "1.3 Chiffrement côté client", 2)
    add_p(
        doc,
        "Le mode Galois/Counter (GCM) associé à AES fournit à la fois la confidentialité et un contrôle d'intégrité (Dworkin, 2007). Une clé de 256 bits et un nonce unique par chiffrement constituent le schéma retenu pour chaque fichier. La clé de fichier ne doit pas être dérivée du seul mot de passe à chaque document : elle est aléatoire, puis enveloppée par une clé maître. Cette clé maître est obtenue par PBKDF2-HMAC-SHA256, fonction de dérivation normalisée pour résister à l'essai exhaustif hors ligne (Moriarty, Kaliski et Rusch, 2017). Le sel est propre au compte et stocké côté serveur ; le mot de passe et la clé maître ne le sont pas.",
    )
    add_p(
        doc,
        "L'OWASP rappelle que le secret de chiffrement des données sensibles ne doit pas résider sur le serveur qui conserve les données, lorsque l'objectif est de limiter l'impact d'une compromission de la base (OWASP Foundation, s. d.). MEGA illustre publiquement un modèle de chiffrement contrôlé par l'utilisateur (Mega Ltd, s. d.). Nextcloud documente un chiffrement de bout en bout distinct du chiffrement serveur (Nextcloud GmbH, s. d.). Ces références confirment la faisabilité du principe, sans imposer leur protocole. ARSAVE en retient l'idée : AES-256-GCM avant transmission, clé maître locale, intégrité du blob par SHA-256 (NIST, 2015).",
    )
    add_h(doc, "1.4 Clients multiplateformes et API", 2)
    add_p(
        doc,
        "Flutter produit, à partir d'une même base, un client mobile et un client web (Flutter team, s. d.). Le client web d'ARSAVE est donc la partie visible du système : tableau de bord, coffre, historique, paramètres et espace d'administration. L'API PHP n'est pas une interface homme-machine ; elle vérifie le jeton, contrôle l'appartenance des fichiers et écrit le blob. MySQL conserve les comptes, les sessions et les métadonnées (Oracle, s. d. ; PHP Group, s. d.).",
    )
    add_h(doc, "1.5 Cadre de conception : Unified Process", 2)
    add_p(
        doc,
        "Le Unified Process organise le projet en phases (inception, élaboration, construction, transition) et en itérations, avec des cas d'utilisation comme fil conducteur (Jacobson, Booch et Rumbaugh, 1999 ; Kruchten, 2004). La notation des modèles est UML 2 (OMG, 2017 ; Rumbaugh, Jacobson et Booch, 2004). Pour ARSAVE, seuls ces artefacts sont produits : cas d'utilisation, scénarios, diagrammes de séquence, diagramme d'activité et diagramme de classes. Ils sont détaillés au chapitre 3. Le chapitre 4 correspond à la construction : réalisation de l'API, du schéma et du client web.",
    )
    add_h(doc, "1.6 Synthèse partielle", 2)
    add_p(
        doc,
        "La revue retient quatre exigences pour le système proposé : une sauvegarde distante versionnée, un chiffrement effectué avant l'envoi, une API qui ne reçoit jamais la clé maître, et une conception guidée par les cas d'utilisation selon UP. Ces exigences ne sont pas encore confrontées aux produits du marché : c'est l'objet du chapitre 2.",
    )
    from memoire_volume import extra_chapitre_1
    extra_chapitre_1(doc, add_h, add_p, add_table)
    add_h(doc, "Conclusion", 2)
    add_p(
        doc,
        "La littérature et les normes examinées montrent qu'un canal TLS (Rescorla, 2018) ne remplace pas un chiffrement du contenu avant transmission (Dworkin, 2007 ; OWASP Foundation, s. d.-a). Elles montrent aussi qu'une clé dérivée du mot de passe doit utiliser un sel et un coût de calcul explicites (Moriarty, Kaliski et Rusch, 2017 ; Turan et al., 2010 ; Bonneau, Herley, van Oorschot et Stajano, 2012). Enfin, le Unified Process fournit le cadre unique de la suite du mémoire : cas d'utilisation, scénarios et modèles UML, sans changer de méthode en cours de route (Jacobson, Booch et Rumbaugh, 1999 ; Kruchten, 2004 ; Larman, 2004).",
    )
    add_p(
        doc,
        "Le chapitre suivant prend ces exigences comme grille de lecture des systèmes existants. Il ne reprend pas la théorie du chiffrement : il vérifie si Drive, Dropbox, OneDrive, MEGA et Nextcloud couvrent déjà le besoin, et il isole ce qui reste à concevoir.",
    )

    add_h(doc, "Chapitre 2. Analyse de systèmes existants", 1)
    add_h(doc, "Introduction", 2)
    add_p(
        doc,
        "Après avoir fixé les exigences, il faut vérifier qu'elles ne sont pas déjà satisfaites par un service courant. Ce chapitre compare des produits accessibles sur le web, à partir de leurs documentations publiques, et non à partir d'une mesure expérimentale de leurs serveurs. L'objectif est d'identifier un écart précis : ce qu'ARSAVE doit faire et que ces systèmes ne réunissent pas dans un déploiement local maîtrisé.",
    )
    add_p(
        doc,
        "La comparaison reste fonctionnelle et liée à la confidentialité. Elle ne classe pas les produits selon leur part de marché. Les critères viennent de la conclusion du chapitre 1 : chiffrement avant envoi, absence de la clé maître sur le serveur, historique, administration des comptes et accès web (Google, s. d. ; Dropbox, s. d. ; Microsoft, s. d. ; Mega Ltd, s. d. ; Nextcloud GmbH, s. d.).",
    )
    add_h(doc, "2.1 Démarche", 2)
    add_p(
        doc,
        "L'analyse porte sur des services largement diffusés qui couvrent le même besoin fonctionnel : déposer un fichier, le retrouver et le récupérer depuis un autre accès. Les critères, issus du chapitre 1, sont le lieu du chiffrement du contenu, la maîtrise de la clé par l'utilisateur, la présence d'un historique, la possibilité d'administrer des comptes et l'existence d'un accès web. Les sources sont les documentations publiques des éditeurs (Google, s. d. ; Dropbox, s. d. ; Microsoft, s. d. ; Mega Ltd, s. d. ; Nextcloud GmbH, s. d.).",
    )
    add_h(doc, "2.2 Google Drive, Dropbox et OneDrive", 2)
    add_p(
        doc,
        "Google Drive, Dropbox et Microsoft OneDrive offrent une copie distante, la synchronisation, la corbeille et un client web abouti. Leur point commun, au regard du besoin d'ARSAVE, est que l'opérateur peut accéder au contenu pour fournir la prévisualisation, la recherche plein texte et le partage. Le chiffrement au repos protège le disque du centre de données, pas la confidentialité vis-à-vis de la plateforme. L'administration des comptes existe dans les offres d'organisation, mais elle porte sur un annuaire de l'éditeur, non sur un serveur que l'établissement contrôle.",
    )
    add_h(doc, "2.3 MEGA", 2)
    add_p(
        doc,
        "MEGA se distingue par un chiffrement réalisé côté client et par une clé liée au compte de l'utilisateur (Mega Ltd, s. d.). Le service est proche de l'objectif de confidentialité d'ARSAVE. Il reste un service public mutualisé : le modèle de compte, les quotas et l'administration ne sont pas ceux d'une application déployée localement avec une base et une API propres. ARSAVE reprend l'idée du chiffrement avant envoi, avec un schéma explicite (sel PBKDF2, clé de fichier enveloppée, nonce GCM) décrit dans la conception.",
    )
    add_h(doc, "2.4 Nextcloud", 2)
    add_p(
        doc,
        "Nextcloud est un logiciel que l'on peut héberger. Il propose le dépôt de fichiers, les versions, la corbeille, le partage et une administration des utilisateurs. Le chiffrement de bout en bout est une fonction distincte du stockage serveur classique (Nextcloud GmbH, s. d.). Le produit est généraliste. ARSAVE est plus étroit : un coffre personnel chiffré avant transmission, un historique d'opérations et une administration des comptes sans lecture du contenu.",
    )
    add_h(doc, "2.5 Comparaison", 2)
    add_p(doc, "Le tableau 1 résume la comparaison selon les critères retenus.", first_line=False)
    add_table(
        doc,
        ["Critère", "Drive / Dropbox / OneDrive", "MEGA", "Nextcloud", "ARSAVE"],
        [
            ["Accès web", "Oui", "Oui", "Oui", "Oui (Flutter web)"],
            ["Chiffrement avant envoi", "Non", "Oui", "Option séparée", "Oui, systématique"],
            ["Clé maître absente du serveur", "Non", "Oui", "Selon le mode", "Oui"],
            ["Versions et corbeille", "Oui", "Oui", "Oui", "Oui"],
            ["Admin des comptes hébergé", "Offre entreprise", "Limité", "Oui", "Oui"],
            ["Lecture du fichier par le serveur", "Possible", "Non prévue", "Possible hors E2E", "Impossible"],
        ],
    )
    add_p(doc, "Tableau 1. Comparaison des systèmes existants et d'ARSAVE", italic=True, center=True, first_line=False)
    add_h(doc, "2.6 Écarts à combler", 2)
    add_p(
        doc,
        "Aucun des services examinés ne réunit, dans un déploiement local simple, le chiffrement systématique avant envoi, une API qui ignore la clé maître, un client web de coffre et une administration des comptes. Ces écarts définissent le système proposé. Ils sont traduits en cas d'utilisation au chapitre 3, conformément au rôle des use cases dans UP (Jacobson, Booch et Rumbaugh, 1999 ; Cockburn, 2001).",
    )
    from memoire_volume import extra_chapitre_2
    extra_chapitre_2(doc, add_h, add_p, add_table)
    add_h(doc, "Conclusion", 2)
    add_p(
        doc,
        "Drive, Dropbox et OneDrive répondent au dépôt et à la reprise des fichiers, mais le contenu reste dans le périmètre de confiance de l'opérateur. MEGA se rapproche du chiffrement côté client, sans offrir le même cadre d'hébergement et d'administration locale. Nextcloud s'héberge et administre des comptes, mais le chiffrement de bout en bout n'est pas le mode par défaut du stockage (Nextcloud GmbH, s. d.).",
    )
    add_p(
        doc,
        "L'écart retenu est donc le suivant : un coffre web dont chaque fichier est chiffré avant l'envoi, avec une administration des comptes qui ne donne pas accès au clair. Le chapitre 3 transforme cet écart en acteurs, cas d'utilisation, scénarios et modèles UML, toujours selon le Unified Process.",
    )

    add_h(doc, "Chapitre 3. Analyse et conception du système proposé", 1)
    add_h(doc, "Introduction", 2)
    add_p(
        doc,
        "Ce chapitre est la phase d'élaboration du Unified Process (Jacobson, Booch et Rumbaugh, 1999 ; Kruchten, 2004). Le périmètre a été fixé aux chapitres 1 et 2. Il s'agit maintenant de décrire le système proposé de façon vérifiable, avant d'écrire le détail du code. La notation est UML 2 (OMG, 2017 ; Rumbaugh, Jacobson et Booch, 2004 ; Fowler, 2003 ; Booch, Rumbaugh et Jacobson, 2005).",
    )
    add_p(
        doc,
        "Les artefacts produits sont uniquement ceux demandés par la méthode retenue : le diagramme de cas d'utilisation, deux scénarios nominaux réalisés en diagrammes de séquence, le diagramme d'activité du dépôt d'un fichier, et le diagramme de classes du domaine. Les descriptions textuelles complètes et les scénarios d'échec sont reportés en annexe, afin de ne pas surcharger les figures.",
    )
    add_h(doc, "3.1 Démarche Unified Process", 2)
    add_p(
        doc,
        "La conception suit le Unified Process (Jacobson, Booch et Rumbaugh, 1999 ; Kruchten, 2004). En inception, le périmètre est fixé : sauvegarder, restaurer et administrer, sans donner le contenu au serveur. En élaboration, les cas d'utilisation structurants sont décrits, puis réalisés par deux scénarios, un diagramme d'activité et un diagramme de classes. La notation est UML 2 (OMG, 2017). La construction est reportée au chapitre 4.",
    )
    add_h(doc, "3.2 Acteurs", 2)
    add_p(
        doc,
        "Deux acteurs interagissent avec le système. L'utilisateur dépose, consulte et restaure ses fichiers. L'administrateur est un utilisateur doté du rôle admin : il consulte les indicateurs globaux et gère les comptes (rôle et suppression). L'administrateur n'est pas un acteur capable de déchiffrer les fichiers des autres comptes, car la clé maître n'est jamais transmise.",
    )
    add_h(doc, "3.3 Diagramme de cas d'utilisation", 2)
    add_p(
        doc,
        "La figure 1 regroupe les cas d'utilisation du système. S'authentifier est le préalable des cas du coffre. Sauvegarder un fichier inclut Chiffrer le fichier : il n'existe pas de dépôt en clair. Gérer les comptes est réservé à l'administrateur. Consulter le tableau de bord, l'historique et le profil complètent le périmètre du client web.",
    )
    add_fig(doc, figs["uc"], "Figure 1. Diagramme de cas d'utilisation d'ARSAVE")
    add_p(
        doc,
        "Le cas S'authentifier a une extension d'échec : identifiants refusés, message affiché, aucune session créée. Le cas Sauvegarder un fichier a une extension d'échec réseau ou de rejet API : le fichier local n'est pas déclaré sauvegardé et l'erreur est montrée à l'utilisateur. Ces extensions ne changent pas le scénario nominal décrit ci-dessous.",
    )
    add_h(doc, "3.4 Scénario 1 — Authentification", 2)
    add_p(
        doc,
        "Précondition : le compte existe. Déclencheur : l'utilisateur valide le formulaire de connexion. Scénario nominal : (1) l'interface envoie l'email et le mot de passe à l'API ; (2) le service retrouve l'utilisateur et vérifie l'empreinte du mot de passe ; (3) une session est créée, seul le hachage SHA-256 du jeton est stocké ; (4) la réponse contient le jeton, le profil et le sel de dérivation ; (5) le client dérive la clé maître par PBKDF2 et ouvre le tableau de bord. Postcondition : une session valide existe et la clé maître est uniquement en mémoire sur le client. La figure 2 représente cet enchaînement.",
    )
    add_fig(doc, figs["s1"], "Figure 2. Diagramme de séquence — authentification (scénario nominal)")
    add_h(doc, "3.5 Scénario 2 — Sauvegarde d'un fichier", 2)
    add_p(
        doc,
        "Précondition : l'utilisateur est authentifié et la clé maître est disponible. Déclencheur : il choisit un fichier dans le coffre. Scénario nominal : (1) le client génère une clé de fichier aléatoire ; (2) le contenu est chiffré en AES-256-GCM ; (3) la clé de fichier est enveloppée avec la clé maître ; (4) le blob, le nom, le type, les nonces, la clé enveloppée et l'empreinte du ciphertext sont envoyés ; (5) l'API contrôle le jeton et le propriétaire, écrit le blob et enregistre la version ; (6) l'interface confirme la sauvegarde. Postcondition : le serveur possède un ciphertext et des métadonnées, jamais le clair. La figure 3 représente cet enchaînement.",
    )
    add_fig(doc, figs["s2"], "Figure 3. Diagramme de séquence — sauvegarde d'un fichier (scénario nominal)")
    add_h(doc, "3.6 Diagramme d'activité", 2)
    add_p(
        doc,
        "Le diagramme d'activité de la figure 4 détaille le flux du cas Sauvegarder un fichier, y compris les deux décisions du scénario : fichier illisible et rejet de l'enregistrement. Le flux nominal correspond au scénario 2. Les branches d'échec affichent un message et n'enregistrent pas de version réussie. Cette lecture complète les diagrammes de séquence, qui ne montrent que le chemin nominal (Rumbaugh, Jacobson et Booch, 2004).",
    )
    add_fig(doc, figs["act"], "Figure 4. Diagramme d'activité — sauvegarder un fichier")
    add_h(doc, "3.7 Diagramme de classes", 2)
    add_p(
        doc,
        "Le diagramme de classes de la figure 5 fixe le modèle du domaine qui sera réalisé dans MySQL et dans les services. Un utilisateur possède plusieurs sessions, appareils, fichiers et opérations. Un fichier possède plusieurs versions ; la version porte le chemin du blob, le nonce et la clé enveloppée. L'opération journalise le type (sauvegarde, téléchargement, corbeille, restauration, suppression) et le statut. Le rôle de l'utilisateur distingue le cas Gérer les comptes.",
    )
    add_fig(doc, figs["cls"], "Figure 5. Diagramme de classes du domaine")
    add_p(
        doc,
        "Ce modèle suffit à passer à la construction : chaque classe d'analyse correspond à une table ou à un service déjà identifiable (authentification, fichiers, chiffrement). Aucun autre formalisme n'est introduit.",
    )
    from memoire_volume import extra_chapitre_3
    extra_chapitre_3(doc, add_h, add_p, add_table)
    add_h(doc, "Conclusion", 2)
    add_p(
        doc,
        "La conception fixe deux acteurs, des cas d'utilisation reliés par une inclusion de chiffrement, deux scénarios nominaux et les branches d'échec du dépôt. Le diagramme de classes relie l'utilisateur à ses sessions, appareils, fichiers, versions et opérations. Ces modèles sont suffisants pour construire le système sans ajouter une autre méthode.",
    )
    add_p(
        doc,
        "Le chapitre 4 montre comment ces artefacts sont réalisés dans le client web, l'API et la base. Les fiches de cas, le détail des tables et les routes HTTP sont placés en annexe pour garder au chapitre la lecture des diagrammes.",
    )

    add_h(doc, "Chapitre 4. Implémentation du système", 1)
    add_h(doc, "Introduction", 2)
    add_p(
        doc,
        "Ce chapitre correspond à la construction dans le Unified Process (Kruchten, 2004 ; Larman, 2004). Il ne redessine pas les cas d'utilisation : il indique comment chacun est porté par un composant du client Flutter, de l'API PHP ou de MySQL. La partie visible du mémoire est le client web, puisque l'application est exécutée dans le navigateur.",
    )
    add_p(
        doc,
        "L'ordre suit les scénarios du chapitre 3. On présente d'abord l'architecture, puis l'authentification, la sauvegarde chiffrée, la restauration, l'historique et l'administration des comptes. Les limites assumées ferment le chapitre. Le schéma relationnel et la liste des routes sont en annexe.",
    )
    add_h(doc, "4.1 Environnement de construction", 2)
    add_p(
        doc,
        "La construction reprend les cas d'utilisation du chapitre 3. Le client est une application Flutter, exécutée dans Chrome pour la partie web et prévue pour Android (Flutter team, s. d.). L'API est écrite en PHP 8 (PHP Group, s. d.) et exposée en REST. Les métadonnées sont dans MySQL, base arsave (Oracle, s. d.). Les blobs sont des fichiers sur disque, sous un répertoire de stockage distinct de la base. La figure 6 situe ces trois niveaux.",
    )
    add_fig(doc, figs["arch"], "Figure 6. Architecture de l'implémentation")
    add_h(doc, "4.2 Réalisation des cas d'utilisation", 2)
    add_p(
        doc,
        "S'inscrire et s'authentifier sont réalisés par AuthController et AuthService. Le mot de passe est conservé sous forme d'empreinte. Le jeton de session est aléatoire ; la base ne stocke que son empreinte SHA-256. Le sel PBKDF2 est renvoyé à la connexion pour que CryptoService dérive la clé maître localement (120 000 itérations, 256 bits), conformément au scénario 1.",
    )
    add_p(
        doc,
        "Sauvegarder un fichier est réalisé par BackupService, qui appelle le chiffrement puis l'envoi multipart du champ blob. FileService écrit le ciphertext, calcule la version et journalise l'opération backup. Restaurer télécharge le blob, vérifie l'empreinte et déchiffre sur le client. La corbeille, la restauration logique et la suppression définitive correspondent aux opérations trash, restore_meta et purge. Le tableau de bord agrège les fichiers et l'historique ; les paramètres permettent le thème sombre ou clair et le rappel de confidentialité.",
    )
    add_p(
        doc,
        "Gérer les comptes est réalisé dans l'écran d'administration du client web, visible seulement si le rôle est admin. L'API /admin/users liste les comptes, change le rôle et supprime un utilisateur après effacement de ses blobs. Le dernier administrateur et le compte courant ne peuvent pas être retirés. Cette réalisation respecte la frontière du cas d'utilisation : l'administrateur ne reçoit aucune clé de déchiffrement.",
    )
    add_h(doc, "4.3 Chiffrement implémenté", 2)
    add_p(
        doc,
        "CryptoService applique AES-256-GCM (Dworkin, 2007). Pour chaque fichier, une clé de 32 octets est tirée, le clair est chiffré, puis la clé est enveloppée avec la clé maître. Le nonce du fichier et le nonce d'enveloppe sont envoyés en base64 avec la clé enveloppée. L'empreinte SHA-256 porte sur le ciphertext (NIST, 2015). Au téléchargement, une empreinte différente bloque le déchiffrement. Le serveur ne voit donc ni le clair, ni la clé maître, ni la clé de fichier (OWASP Foundation, s. d.).",
    )
    add_h(doc, "4.4 Partie web", 2)
    add_p(
        doc,
        "La partie web reprend les mêmes cas que le client : connexion, tableau de bord (indicateurs, répartition, histogramme, historique), coffre avec miniatures pour les images déchiffrées localement, activité, corbeille et paramètres. Les messages d'échec des extensions de cas (identifiants, réseau, aperçu) sont affichés à l'utilisateur. Le sélecteur de fichiers accepte tout type de contenu ; la limite pratique d'envoi est celle configurée sur l'API.",
    )
    add_h(doc, "4.5 Limites de la construction", 2)
    add_p(
        doc,
        "La réinitialisation du mot de passe régénère le sel : les clés enveloppées existantes ne sont plus ouvrables sans une procédure de re-chiffrement, qui n'est pas dans le périmètre réalisé. Les miniatures ne sont produites que pour les images de taille limitée. L'administration ne gère pas encore la création d'un compte ni la réinitialisation d'un mot de passe depuis l'écran admin. Ces limites sont des travaux de transition au sens de UP, non des fonctions contradictoires avec les scénarios nominaux.",
    )
    from memoire_volume import extra_chapitre_4
    extra_chapitre_4(doc, add_h, add_p, add_table)
    add_h(doc, "Conclusion", 2)
    add_p(
        doc,
        "L'implémentation couvre les cas d'utilisation élaborés au chapitre 3 : authentification, sauvegarde chiffrée, restauration, historique, corbeille et administration des comptes sur le client web. Le serveur conserve des métadonnées et des blobs chiffrés. La clé maître reste sur le poste de l'utilisateur, comme l'exigeaient la revue de littérature et l'écart identifié face aux systèmes existants.",
    )
    add_p(
        doc,
        "La construction reste fidèle aux scénarios nominaux. Les extensions d'échec affichent un message plutôt que d'enregistrer une sauvegarde incomplète. Les travaux laissés à une itération ultérieure — re-chiffrement après changement de mot de passe, création de compte depuis l'écran admin — sont des compléments, non une remise en cause du modèle de classes ni du chiffrement AES-256-GCM (Dworkin, 2007 ; NIST, 2001).",
    )

    add_h(doc, "Bibliographie", 1)
    add_p(
        doc,
        "Les références sont classées par ordre alphabétique du premier auteur. Les documents sans date sont des documentations officielles consultées pour le mémoire.",
        first_line=False,
    )
    refs = [
        "Anderson, R. (2020). Security engineering: A guide to building dependable distributed systems (3e éd.). Wiley.",
        "Booch, G., Rumbaugh, J., & Jacobson, I. (2005). The Unified Modeling Language user guide (2e éd.). Addison-Wesley.",
        "Bonneau, J., Herley, C., van Oorschot, P. C., & Stajano, F. (2012). The quest to replace passwords: A framework for comparative evaluation of web authentication schemes. IEEE Symposium on Security and Privacy, 553–567.",
        "Bray, T. (2017). The JavaScript Object Notation (JSON) data interchange format (RFC 8259). IETF.",
        "Cockburn, A. (2001). Writing effective use cases. Addison-Wesley.",
        "Daemen, J., & Rijmen, V. (2002). The design of Rijndael: AES — The Advanced Encryption Standard. Springer.",
        "Dropbox. (s. d.). Dropbox Business security whitepaper. Documentation publique Dropbox.",
        "Dworkin, M. (2007). Recommendation for block cipher modes of operation: Galois/Counter Mode (GCM) and GMAC (NIST Special Publication 800-38D). National Institute of Standards and Technology.",
        "Eastlake, D., & Hansen, T. (2011). US secure hash algorithms (RFC 6234). IETF.",
        "Ferguson, N., Schneier, B., & Kohno, T. (2010). Cryptography engineering: Design principles and practical applications. Wiley.",
        "Fielding, R. T. (2000). Architectural styles and the design of network-based software architectures [Thèse de doctorat, University of California, Irvine].",
        "Fielding, R., & Reschke, J. (2014). Hypertext Transfer Protocol (HTTP/1.1): Semantics and content (RFC 7231). IETF.",
        "Florêncio, D., & Herley, C. (2007). A large-scale study of web password habits. Proceedings of the 16th International Conference on World Wide Web, 657–666.",
        "Flutter team. (s. d.). Flutter documentation. https://docs.flutter.dev",
        "Fowler, M. (2003). UML distilled: A brief guide to the standard object modeling language (3e éd.). Addison-Wesley.",
        "Google. (s. d.). Google Workspace security whitepaper. Documentation publique Google.",
        "Grassi, P. A., Garcia, M. E., & Fenton, J. L. (2017). Digital identity guidelines (NIST Special Publication 800-63-3). National Institute of Standards and Technology.",
        "ISO/IEC. (2011). ISO/IEC 25010:2011 — Systems and software Quality Requirements and Evaluation (SQuaRE) — System and software quality models.",
        "ISO/IEC. (2022). ISO/IEC 27001:2022 — Information security, cybersecurity and privacy protection — Information security management systems — Requirements.",
        "Jacobson, I., Booch, G., & Rumbaugh, J. (1999). The Unified Software Development Process. Addison-Wesley.",
        "Josefsson, S. (2006). The Base16, Base32, and Base64 data encodings (RFC 4648). IETF.",
        "Katz, J., & Lindell, Y. (2020). Introduction to modern cryptography (3e éd.). CRC Press.",
        "Krawczyk, H., Bellare, M., & Canetti, R. (1997). HMAC: Keyed-hashing for message authentication (RFC 2104). IETF.",
        "Kruchten, P. (2004). The Rational Unified Process: An introduction (3e éd.). Addison-Wesley.",
        "Larman, C. (2004). Applying UML and patterns: An introduction to object-oriented analysis and design and iterative development (3e éd.). Prentice Hall.",
        "Leach, P., Mealling, M., & Salz, R. (2005). A universally unique identifier (UUID) URN namespace (RFC 4122). IETF.",
        "Masinter, L. (2015). Returning values from forms: multipart/form-data (RFC 7578). IETF.",
        "Mega Ltd. (s. d.). Security and privacy. https://mega.io/security",
        "Menezes, A. J., van Oorschot, P. C., & Vanstone, S. A. (1996). Handbook of applied cryptography. CRC Press.",
        "Microsoft. (s. d.). OneDrive and Microsoft 365 security documentation. Documentation publique Microsoft.",
        "Moriarty, K., Kaliski, B., & Rusch, A. (2017). PKCS #5: Password-based cryptography specification version 2.1 (RFC 8018). IETF.",
        "Nextcloud GmbH. (s. d.). Nextcloud documentation: encryption. https://docs.nextcloud.com",
        "NIST. (2001). Advanced Encryption Standard (FIPS 197). National Institute of Standards and Technology.",
        "NIST. (2015). Secure hash standard (FIPS 180-4). National Institute of Standards and Technology.",
        "Object Management Group. (2017). OMG Unified Modeling Language (OMG UML), version 2.5.1.",
        "Oracle. (s. d.). MySQL 8.0 reference manual. https://dev.mysql.com/doc/",
        "OWASP Foundation. (s. d.-a). Cryptographic storage cheat sheet. https://cheatsheetseries.owasp.org",
        "OWASP Foundation. (s. d.-b). Authentication cheat sheet. https://cheatsheetseries.owasp.org",
        "OWASP Foundation. (s. d.-c). Session management cheat sheet. https://cheatsheetseries.owasp.org",
        "PHP Group. (s. d.). PHP manual. https://www.php.net/manual/en/",
        "Pressman, R. S., & Maxim, B. R. (2019). Software engineering: A practitioner's approach (9e éd.). McGraw-Hill.",
        "Rescorla, E. (2018). The Transport Layer Security (TLS) protocol version 1.3 (RFC 8446). IETF.",
        "Rumbaugh, J., Jacobson, I., & Booch, G. (2004). The Unified Modeling Language reference manual (2e éd.). Addison-Wesley.",
        "Saltzer, J. H., & Schroeder, M. D. (1975). The protection of information in computer systems. Proceedings of the IEEE, 63(9), 1278–1308.",
        "Sommerville, I. (2016). Software engineering (10e éd.). Pearson.",
        "Stallings, W. (2020). Cryptography and network security: Principles and practice (8e éd.). Pearson.",
        "Turan, M. S., Barker, E., Burr, W., & Chen, L. (2010). Recommendation for password-based key derivation. Part 1: Storage applications (NIST Special Publication 800-132). National Institute of Standards and Technology.",
    ]
    for ref in refs:
        add_p(doc, ref, size=11, space_after=4, first_line=False)

    add_h(doc, "Annexes", 1)

    add_h(doc, "Annexe A. Fiches des cas d'utilisation", 2)
    add_p(
        doc,
        "Les fiches complètent la figure 1. Seuls les cas structurants sont détaillés. Les autres cas (tableau de bord, historique, profil, corbeille) suivent le même schéma : acteur authentifié, précondition de session, postcondition mise à jour de l'affichage.",
        first_line=False,
    )
    add_table(
        doc,
        ["Cas", "Acteur", "Précondition", "Nominal", "Échec"],
        [
            ["S'inscrire", "Utilisateur", "Email libre", "Compte créé, sel généré", "Email déjà utilisé"],
            ["S'authentifier", "Utilisateur", "Compte existant", "Jeton + sel, clé maître locale", "Identifiants refusés"],
            ["Sauvegarder", "Utilisateur", "Session et clé maître", "Blob chiffré et version créée", "Fichier illisible ou API en échec"],
            ["Restaurer", "Utilisateur", "Fichier possédé", "Clair restitué sur le client", "Empreinte invalide"],
            ["Gérer les comptes", "Administrateur", "Rôle admin", "Rôle modifié ou compte supprimé", "Dernier admin ou soi-même"],
        ],
    )
    add_p(doc, "Tableau A1. Fiches résumées des cas d'utilisation", italic=True, center=True, first_line=False)

    add_h(doc, "Annexe B. Scénarios alternatifs", 2)
    add_p(
        doc,
        "Scénario alternatif A1, authentification refusée. L'utilisateur soumet un email ou un mot de passe incorrect. L'API compare l'empreinte, ne crée pas de session et répond par une erreur d'identifiants. Le client affiche le message et ne dérive pas de clé maître. Aucune donnée de coffre n'est accessible.",
        first_line=False,
    )
    add_p(
        doc,
        "Scénario alternatif A2, sauvegarde interrompue. Le fichier est choisi mais la lecture locale échoue, ou l'API refuse l'envoi (session absente, champ manquant, limite de taille). Le client affiche l'erreur. Aucune version réussie n'est journalisée comme un succès. Le flux correspond aux branches « non » de la figure 4.",
    )

    add_h(doc, "Annexe C. Tables réalisées", 2)
    add_p(
        doc,
        "Le schéma MySQL reprend le diagramme de classes. Les colonnes cryptographiques ne contiennent jamais la clé maître.",
        first_line=False,
    )
    add_table(
        doc,
        ["Table", "Rôle", "Colonnes principales"],
        [
            ["users", "Comptes", "email, password_hash, name, role, kdf_salt"],
            ["devices", "Appareils", "user_id, device_uid, device_name"],
            ["sessions", "Sessions", "user_id, token_hash, expires_at"],
            ["files", "Fichiers", "user_id, logical_name, mime, status"],
            ["file_versions", "Versions", "nonce_b64, wrapped_key_b64, checksum_sha256, storage_path"],
            ["operations", "Historique", "type, status, message, created_at"],
            ["password_resets", "Réinitialisation", "token_hash, expires_at, used_at"],
        ],
    )
    add_p(doc, "Tableau C1. Tables de la base arsave", italic=True, center=True, first_line=False)

    add_h(doc, "Annexe D. Routes de l'API", 2)
    add_table(
        doc,
        ["Méthode", "Route", "Cas d'utilisation"],
        [
            ["POST", "/auth/register", "S'inscrire"],
            ["POST", "/auth/login", "S'authentifier"],
            ["POST", "/auth/logout", "Quitter la session"],
            ["GET", "/files", "Consulter le coffre ou la corbeille"],
            ["POST", "/files", "Sauvegarder un fichier"],
            ["GET", "/files/{id}/download", "Restaurer un fichier"],
            ["POST", "/files/{id}/trash", "Mettre en corbeille"],
            ["POST", "/files/{id}/restore", "Remettre dans le coffre"],
            ["POST", "/files/{id}/purge", "Supprimer définitivement"],
            ["GET", "/operations", "Consulter l'historique"],
            ["GET", "/admin/users", "Gérer les comptes"],
            ["PUT", "/admin/users/{id}", "Changer le rôle ou le nom"],
            ["DELETE", "/admin/users/{id}", "Supprimer un compte"],
        ],
    )
    add_p(doc, "Tableau D1. Principales routes reliées aux cas d'utilisation", italic=True, center=True, first_line=False)

    add_h(doc, "Annexe E. Glossaire", 2)
    add_table(
        doc,
        ["Terme", "Sens dans ARSAVE"],
        [
            ["UP", "Unified Process : méthode de conduite retenue pour tout le mémoire"],
            ["UML", "Notation des cas, séquences, activités et classes"],
            ["Clé maître", "Clé dérivée du mot de passe par PBKDF2, présente seulement sur le client"],
            ["Clé de fichier", "Clé aléatoire AES-256, enveloppée par la clé maître"],
            ["Nonce", "Valeur unique associée à un chiffrement GCM"],
            ["Blob", "Fichier déjà chiffré, stocké sur disque"],
            ["Jeton", "Secret de session ; seule son empreinte SHA-256 est en base"],
            ["Acteur", "Utilisateur ou administrateur, au sens des cas d'utilisation"],
        ],
    )
    add_p(doc, "Tableau E1. Glossaire", italic=True, center=True, first_line=False)

    from memoire_volume import extra_annexes
    extra_annexes(doc, add_h, add_p, add_table)
    add_h(doc, "Annexe F. Jeu d'essai local", 2)
    add_p(
        doc,
        "Le jeu d'essai sert à vérifier les scénarios sur la partie web. Un compte administrateur de démonstration est chargé au démarrage local : admin@arsave.local. Le mot de passe de démonstration est celui du script de graine du projet et doit être changé hors d'un poste de développement. Un second compte créé par le formulaire d'inscription permet de vérifier qu'un utilisateur simple ne voit pas l'entrée d'administration.",
        first_line=False,
    )
    add_p(
        doc,
        "Contrôles minimaux : connexion refusée si le mot de passe est faux (annexe B) ; sauvegarde d'un petit fichier puis présence d'une ligne d'historique ; restauration qui déclenche le téléchargement du clair ; suppression d'un compte de test depuis l'administrateur, sans effet sur le dernier compte admin.",
    )

    out = OUT / "Memoire_ARSAVE_chapitres_1_a_4_complet.docx"
    try:
        doc.save(out)
    except PermissionError:
        out = OUT / "Memoire_ARSAVE.docx"
        doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
